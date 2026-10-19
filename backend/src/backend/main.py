"""
Khmer DocFixer: PDF-to-PPTX Unicode Restorer - FastAPI Backend.
Main application server with PPTX parsing, PDF alignment, Khmer Unicode restoration,
and presentation streaming endpoints.
"""

import logging
import os
import shutil
import uuid
from typing import Dict, List, Optional
from fastapi import FastAPI, File, Form, HTTPException, UploadFile, BackgroundTasks
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse
import docx
from pptx import Presentation

from .models import (
    ApplyFixesRequest,
    FontOption,
    ProcessResponse,
    ReplacementItem,
    SlideDiff,
)
from .engine.dictionary import (
    replace_dictionary_terms,
    repair_khmer_heuristics,
    clean_khmer_unicode,
    restore_khmer_text,
    CUSTOM_DICT_FILE,
    get_user_custom_dictionary,
)
from .engine.legacy_converter import convert_limon_to_unicode, is_likely_legacy_font
from .engine.pdf_extractor import extract_text_from_pdf
from .engine.pdf_to_pptx import convert_pdf_to_pptx
from .engine.pdf_processor import apply_pdf_replacements_and_fonts
from .engine.office_exporter import export_pptx_to_pdf, export_docx_to_pdf
from .engine.text_aligner import align_slide_with_pdf
from .engine.gemini_restorer import restore_with_gemini
from .engine.pptx_processor import (
    RECOMMENDED_KHMER_FONTS,
    apply_replacements_and_fonts,
    extract_presentation_runs,
)
from .engine.docx_processor import (
    extract_docx_sections_or_pages,
    apply_docx_replacements_and_fonts,
)
from .samples.sample_generator import generate_sample_files
from pydantic import BaseModel

logging.basicConfig(level=logging.INFO)
logger = logging.getLogger("khmer_docfixer")

app = FastAPI(
    title="Khmer DocFixer API",
    description="Backend API for restoring corrupted Khmer legacy text in PowerPoint decks using reference PDFs, dictionaries, and LLM.",
    version="1.0.0"
)

# CORS Middleware
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

import tempfile
import json

BASE_TEMP_DIR = os.environ.get("TEMP_DIR") or os.path.join(tempfile.gettempdir(), "khmer_docfixer_sessions")
try:
    os.makedirs(BASE_TEMP_DIR, exist_ok=True)
except Exception as e:
    logger.warning(f"Could not initialize BASE_TEMP_DIR at {BASE_TEMP_DIR}: {e}")

# Session metadata storage with disk fallback for serverless container reloads
session_registry: Dict[str, Dict] = {}


def save_session_meta(session_id: str, meta: dict):
    """Saves session metadata in-memory and persists meta.json to disk."""
    session_registry[session_id] = meta
    try:
        session_dir = os.path.join(BASE_TEMP_DIR, session_id)
        os.makedirs(session_dir, exist_ok=True)
        meta_file = os.path.join(session_dir, "meta.json")
        with open(meta_file, "w", encoding="utf-8") as f:
            json.dump(meta, f)
    except Exception as e:
        logger.warning(f"Could not persist session meta for {session_id}: {e}")


def get_session_meta(session_id: str) -> Optional[dict]:
    """Retrieves session metadata from memory or disk cache."""
    if session_id in session_registry:
        return session_registry[session_id]
    meta_file = os.path.join(BASE_TEMP_DIR, session_id, "meta.json")
    if os.path.exists(meta_file):
        try:
            with open(meta_file, "r", encoding="utf-8") as f:
                meta = json.load(f)
                session_registry[session_id] = meta
                return meta
        except Exception as e:
            logger.warning(f"Failed to read session meta for {session_id}: {e}")
    return None


def cleanup_session_dir(session_id: str):
    """Deletes temporary session directory and cached data."""
    try:
        session_dir = os.path.join(BASE_TEMP_DIR, session_id)
        if os.path.exists(session_dir):
            shutil.rmtree(session_dir, ignore_errors=True)
        session_registry.pop(session_id, None)
        logger.info(f"Cleaned up session {session_id}")
    except Exception as e:
        logger.error(f"Error cleaning session {session_id}: {e}")


@app.get("/api/health")
@app.get("/health")
def health_check():
    return {"status": "ok", "app": "Khmer DocFixer: PDF-to-PPTX Unicode Restorer"}


@app.get("/api/fonts", response_model=List[FontOption])
@app.get("/fonts", response_model=List[FontOption])
def get_fonts():
    return [FontOption(**f) for f in RECOMMENDED_KHMER_FONTS]


class CustomDictEntryRequest(BaseModel):
    wrong: Optional[str] = None
    right: Optional[str] = None
    raw_lines: Optional[str] = None
    replace_all: bool = False


@app.get("/api/custom-dictionary")
@app.get("/custom-dictionary")
def get_custom_dictionary():
    """Returns the user's custom dictionary entries and raw file content."""
    entries_dict = get_user_custom_dictionary()
    raw_content = ""
    target_file = CUSTOM_DICT_FILE
    fallback_file = os.path.join(tempfile.gettempdir(), "custom_dictionary.txt")
    if os.path.exists(target_file):
        try:
            with open(target_file, "r", encoding="utf-8") as f:
                raw_content = f.read()
        except Exception:
            pass
    elif os.path.exists(fallback_file):
        try:
            with open(fallback_file, "r", encoding="utf-8") as f:
                raw_content = f.read()
        except Exception:
            pass
    entries = [{"wrong": k, "right": v} for k, v in entries_dict.items()]
    return {
        "entries": entries,
        "count": len(entries),
        "raw_content": raw_content,
        "file_path": target_file,
    }


@app.post("/api/custom-dictionary")
@app.post("/custom-dictionary")
def update_custom_dictionary(payload: CustomDictEntryRequest):
    """Adds a new wrong -> right pair, appends bulk lines, or saves the full custom_dictionary.txt."""
    lines_to_append = []
    if payload.wrong and payload.right:
        lines_to_append.append(f"{payload.wrong.strip()} = {payload.right.strip()}")
    if payload.raw_lines and not payload.replace_all:
        for line in payload.raw_lines.splitlines():
            if line.strip():
                lines_to_append.append(line.strip())

    for target_path in [CUSTOM_DICT_FILE, os.path.join(tempfile.gettempdir(), "custom_dictionary.txt")]:
        try:
            os.makedirs(os.path.dirname(target_path), exist_ok=True)
            if payload.replace_all and payload.raw_lines is not None:
                with open(target_path, "w", encoding="utf-8") as f:
                    f.write(payload.raw_lines.strip() + "\n")
            elif lines_to_append:
                existing = ""
                if os.path.exists(target_path):
                    with open(target_path, "r", encoding="utf-8") as f:
                        existing = f.read()
                with open(target_path, "a", encoding="utf-8") as f:
                    if existing and not existing.endswith("\n"):
                        f.write("\n")
                    for ln in lines_to_append:
                        f.write(ln + "\n")
            break
        except (OSError, PermissionError):
            continue

    entries_dict = get_user_custom_dictionary()
    raw_content = ""
    for check_path in [CUSTOM_DICT_FILE, os.path.join(tempfile.gettempdir(), "custom_dictionary.txt")]:
        if os.path.exists(check_path):
            try:
                with open(check_path, "r", encoding="utf-8") as f:
                    raw_content = f.read()
                break
            except Exception:
                pass
    entries = [{"wrong": k, "right": v} for k, v in entries_dict.items()]
    return {
        "status": "ok",
        "entries": entries,
        "count": len(entries),
        "raw_content": raw_content,
    }


@app.post("/api/custom-dictionary/upload")
@app.post("/custom-dictionary/upload")
async def upload_custom_dictionary_file(file: UploadFile = File(...)):
    """Uploads a .txt file of custom dictionary entries and appends them to custom_dictionary.txt."""
    content_bytes = await file.read()
    try:
        text_content = content_bytes.decode("utf-8-sig")
    except UnicodeDecodeError:
        text_content = content_bytes.decode("utf-8", errors="ignore")

    lines = [ln.strip() for ln in text_content.splitlines() if ln.strip() and not ln.strip().startswith("#")]
    if lines:
        for target_path in [CUSTOM_DICT_FILE, os.path.join(tempfile.gettempdir(), "custom_dictionary.txt")]:
            try:
                os.makedirs(os.path.dirname(target_path), exist_ok=True)
                existing = ""
                if os.path.exists(target_path):
                    with open(target_path, "r", encoding="utf-8") as f:
                        existing = f.read()
                with open(target_path, "a", encoding="utf-8") as f:
                    if existing and not existing.endswith("\n"):
                        f.write("\n")
                    for ln in lines:
                        f.write(ln + "\n")
                break
            except (OSError, PermissionError):
                continue

    entries_dict = get_user_custom_dictionary()
    raw_content = ""
    for check_path in [CUSTOM_DICT_FILE, os.path.join(tempfile.gettempdir(), "custom_dictionary.txt")]:
        if os.path.exists(check_path):
            try:
                with open(check_path, "r", encoding="utf-8") as f:
                    raw_content = f.read()
                break
            except Exception:
                pass
    entries = [{"wrong": k, "right": v} for k, v in entries_dict.items()]
    return {
        "status": "ok",
        "added_lines": len(lines),
        "entries": entries,
        "count": len(entries),
        "raw_content": raw_content,
    }



def _process_presentation_internal(
    session_id: str,
    pptx_path: str,
    pdf_path: Optional[str] = None,
    mode: str = "auto",
    target_font: str = "Khmer OS Battambang",
    gemini_api_key: Optional[str] = None
) -> ProcessResponse:
    """Core analysis and detection engine."""
    prs = Presentation(pptx_path)
    runs = extract_presentation_runs(prs)
    
    # 1. Extract PDF pages if PDF provided
    pdf_pages = []
    if pdf_path and os.path.exists(pdf_path):
        pdf_pages = extract_text_from_pdf(pdf_path)
        logger.info(f"Extracted {len(pdf_pages)} pages from reference PDF.")

    all_replacements: List[ReplacementItem] = []
    slides_diff: List[SlideDiff] = []

    # Process each slide
    for slide_idx, slide in enumerate(prs.slides):
        slide_runs = [r for r in runs if r.slide_index == slide_idx]
        slide_title = ""
        if slide.shapes.title and slide.shapes.title.has_text_frame:
            slide_title = slide.shapes.title.text_frame.text.strip()
            
        full_slide_text_pieces = []
        slide_replacements: List[ReplacementItem] = []
        seen_replacement_pairs = set()

        # Step A: Slide-to-PDF Alignment (if PDF available)
        pdf_aligned_rules = []
        if pdf_pages and mode in ("auto", "pdf_alignment"):
            runs_dict = [r.to_dict() for r in slide_runs]
            pdf_aligned_rules = align_slide_with_pdf(slide_idx, runs_dict, pdf_pages)
            for rule in pdf_aligned_rules:
                orig = rule["original"].strip()
                repl = rule["replacement"].strip()
                pair_key = (slide_idx, orig, repl)
                if orig and repl and orig != repl and pair_key not in seen_replacement_pairs:
                    seen_replacement_pairs.add(pair_key)
                    item_id = str(uuid.uuid4())[:8]
                    rep_item = ReplacementItem(
                        id=item_id,
                        slide_index=slide_idx,
                        shape_id=rule.get("shape_id", ""),
                        paragraph_index=rule.get("paragraph_index"),
                        run_index=rule.get("run_index"),
                        original=orig,
                        replacement=repl,
                        confidence=rule.get("confidence", 0.92),
                        source="pdf_alignment",
                        status="accepted",
                        explanation="Aligned with reference PDF content",
                        context=rule.get("context")
                    )
                    slide_replacements.append(rep_item)
                    all_replacements.append(rep_item)

        # Step B: Limon / ABC Legacy font conversion and Dictionary/Heuristics
        for para_item in slide_runs:
            raw_text = para_item.text.strip()
            if not raw_text:
                continue
            full_slide_text_pieces.append(raw_text)

            # Check if this paragraph is legacy Limon/ABC
            if is_likely_legacy_font(para_item.font_name, raw_text):
                converted = convert_limon_to_unicode(raw_text)
                if converted.strip() and converted.strip() != raw_text:
                    pair_key = (slide_idx, raw_text, converted.strip())
                    if pair_key not in seen_replacement_pairs:
                        seen_replacement_pairs.add(pair_key)
                        item_id = str(uuid.uuid4())[:8]
                        rep_item = ReplacementItem(
                            id=item_id,
                            slide_index=slide_idx,
                            shape_id=para_item.shape_id,
                            paragraph_index=para_item.paragraph_index,
                            run_index=0,
                            original=raw_text,
                            replacement=converted.strip(),
                            confidence=0.95,
                            source="limon_translit",
                            status="accepted",
                            explanation=f"Legacy font transliteration ({para_item.font_name or 'Limon'}) to Khmer Unicode",
                            context=raw_text
                        )
                        slide_replacements.append(rep_item)
                        all_replacements.append(rep_item)

            # Check Dictionary & Heuristic Replacements on the unified paragraph
            fixed_para, matches = restore_khmer_text(raw_text)
            for orig_term, repl_term in matches:
                pair_key = (slide_idx, orig_term, repl_term)
                if pair_key not in seen_replacement_pairs:
                    seen_replacement_pairs.add(pair_key)
                    item_id = str(uuid.uuid4())[:8]
                    rep_item = ReplacementItem(
                        id=item_id,
                        slide_index=slide_idx,
                        shape_id=para_item.shape_id,
                        paragraph_index=para_item.paragraph_index,
                        run_index=0,
                        original=orig_term,
                        replacement=repl_term,
                        confidence=0.99,
                        source="dictionary",
                        status="accepted",
                        explanation="Matched in Khmer corruption dictionary",
                        context=raw_text
                    )
                    slide_replacements.append(rep_item)
                    all_replacements.append(rep_item)

            # If the paragraph changed as a whole without a sub-match
            if fixed_para != raw_text and not matches:
                pair_key = (slide_idx, raw_text, fixed_para)
                if pair_key not in seen_replacement_pairs:
                    seen_replacement_pairs.add(pair_key)
                    item_id = str(uuid.uuid4())[:8]
                    rep_item = ReplacementItem(
                        id=item_id,
                        slide_index=slide_idx,
                        shape_id=para_item.shape_id,
                        paragraph_index=para_item.paragraph_index,
                        run_index=0,
                        original=raw_text,
                        replacement=fixed_para,
                        confidence=0.92,
                        source="heuristic",
                        status="accepted",
                        explanation="Khmer Unicode heuristic normalization & coeng repair",
                        context=raw_text
                    )
                    slide_replacements.append(rep_item)
                    all_replacements.append(rep_item)

        # Step C: Optional Gemini AI Restoration
        if (mode == "gemini_ai" or gemini_api_key) and (gemini_api_key or os.getenv("GEMINI_API_KEY")):
            pdf_ctx = ""
            if pdf_pages and slide_idx < len(pdf_pages):
                pdf_ctx = pdf_pages[slide_idx].text
            
            gemini_runs = [r.to_dict() for r in slide_runs]
            gemini_results = restore_with_gemini(
                slide_index=slide_idx,
                texts_to_fix=gemini_runs,
                pdf_context=pdf_ctx,
                api_key=gemini_api_key
            )
            for g_item in gemini_results:
                orig = g_item["original"].strip()
                repl = g_item["replacement"].strip()
                pair_key = (slide_idx, orig, repl)
                if orig and repl and orig != repl and pair_key not in seen_replacement_pairs:
                    seen_replacement_pairs.add(pair_key)
                    item_id = str(uuid.uuid4())[:8]
                    rep_item = ReplacementItem(
                        id=item_id,
                        slide_index=slide_idx,
                        shape_id=g_item.get("shape_id", ""),
                        paragraph_index=g_item.get("paragraph_index"),
                        run_index=g_item.get("run_index"),
                        original=orig,
                        replacement=repl,
                        confidence=g_item.get("confidence", 0.98),
                        source="gemini_ai",
                        status="accepted",
                        explanation=g_item.get("explanation", "Gemini AI intelligent restoration"),
                        context=g_item.get("original")
                    )
                    slide_replacements.append(rep_item)
                    all_replacements.append(rep_item)

        # Compute slide original vs preview corrected text
        original_text = "\n".join(full_slide_text_pieces)
        corrected_paragraphs = []
        for piece in full_slide_text_pieces:
            fixed_p, _ = restore_khmer_text(piece)
            corrected_paragraphs.append(fixed_p)
        corrected_text = "\n".join(corrected_paragraphs)

        slides_diff.append(SlideDiff(
            slide_index=slide_idx,
            slide_number=slide_idx + 1,
            title=slide_title or f"Slide {slide_idx + 1}",
            original_text=original_text,
            preview_corrected_text=corrected_text,
            replacements=slide_replacements,
            shape_count=len(slide.shapes),
            table_count=sum(1 for s in slide.shapes if s.has_table)
        ))

    return ProcessResponse(
        session_id=session_id,
        total_slides=len(prs.slides),
        total_corrupted_found=len(all_replacements),
        slides=slides_diff,
        all_replacements=all_replacements,
        target_font=target_font,
        has_pdf_reference=bool(pdf_pages),
        mode=mode
    )


@app.post("/api/upload", response_model=ProcessResponse)
@app.post("/upload", response_model=ProcessResponse)
async def upload_and_process(
    file: Optional[UploadFile] = File(None),
    pptx_file: Optional[UploadFile] = File(None),
    pdf_file: Optional[UploadFile] = File(None),
    target_font: str = Form("Khmer OS Battambang"),
    mode: str = Form("auto"),
    gemini_api_key: Optional[str] = Form(None)
):
    """
    Accepts a PowerPoint (.pptx), Word (.docx), or PDF (.pdf) file.
    - If PDF: performs Khmer OCR and spelling check to convert into an editable PowerPoint (.pptx).
    - If PPTX/DOCX: analyzes all text, performs Khmer spelling check & dictionary-based restoration.
    """
    upload_doc = file or pptx_file
    if not upload_doc or not upload_doc.filename:
        if pdf_file and pdf_file.filename:
            upload_doc = pdf_file
            pdf_file = None
        else:
            raise HTTPException(status_code=400, detail="Please upload a PowerPoint (.pptx), Word (.docx), or PDF (.pdf) file.")

    fname_lower = upload_doc.filename.lower()
    if not (fname_lower.endswith(".pptx") or fname_lower.endswith(".docx") or fname_lower.endswith(".pdf")):
        raise HTTPException(
            status_code=400,
            detail="Unsupported format. Please upload a PowerPoint (.pptx), Word (.docx), or PDF (.pdf) document."
        )

    is_docx = fname_lower.endswith(".docx")
    is_pdf = fname_lower.endswith(".pdf")
    doc_type = "docx" if is_docx else ("pdf" if is_pdf else "pptx")

    session_id = str(uuid.uuid4())
    session_dir = os.path.join(BASE_TEMP_DIR, session_id)
    os.makedirs(session_dir, exist_ok=True)

    saved_doc_name = "input.docx" if is_docx else ("input.pdf" if is_pdf else "input.pptx")
    doc_save_path = os.path.join(session_dir, saved_doc_name)
    content = await upload_doc.read()
    if not content or len(content) == 0:
        raise HTTPException(
            status_code=400,
            detail="The uploaded file is empty (0 bytes). Please upload a valid document."
        )
    with open(doc_save_path, "wb") as f:
        f.write(content)

    pdf_save_path = None
    if pdf_file and pdf_file.filename and pdf_file.filename.lower().endswith(".pdf"):
        pdf_save_path = os.path.join(session_dir, "reference.pdf")
        with open(pdf_save_path, "wb") as f:
            pdf_content = await pdf_file.read()
            f.write(pdf_content)

    session_meta = {
        "doc_path": doc_save_path,
        "pptx_path": doc_save_path if not (is_docx or is_pdf) else None,
        "docx_path": doc_save_path if is_docx else None,
        "pdf_path": doc_save_path if is_pdf else pdf_save_path,
        "original_filename": upload_doc.filename,
        "target_font": target_font,
        "mode": mode,
        "gemini_api_key": gemini_api_key,
        "document_type": doc_type
    }
    save_session_meta(session_id, session_meta)

    try:
        if is_pdf:
            output_pptx_path = os.path.join(session_dir, "converted_presentation.pptx")
            force_ocr = (mode == "ocr")
            converted_path, slides_diff, all_replacements = convert_pdf_to_pptx(
                pdf_path=doc_save_path,
                output_pptx_path=output_pptx_path,
                temp_dir=session_dir,
                target_font=target_font,
                force_ocr=force_ocr
            )
            session_meta["pptx_path"] = converted_path
            save_session_meta(session_id, session_meta)
            return ProcessResponse(
                session_id=session_id,
                total_slides=len(slides_diff),
                total_corrupted_found=len(all_replacements),
                slides=slides_diff,
                all_replacements=all_replacements,
                target_font=target_font,
                has_pdf_reference=True,
                mode=mode,
                document_type="pdf",
                filename=upload_doc.filename
            )
        elif is_docx:
            slides_diff, all_replacements = extract_docx_sections_or_pages(
                docx_path=doc_save_path,
                target_font=target_font
            )
            return ProcessResponse(
                session_id=session_id,
                total_slides=len(slides_diff),
                total_corrupted_found=len(all_replacements),
                slides=slides_diff,
                all_replacements=all_replacements,
                target_font=target_font,
                has_pdf_reference=False,
                mode=mode,
                document_type="docx",
                filename=upload_doc.filename
            )
        else:
            response = _process_presentation_internal(
                session_id=session_id,
                pptx_path=doc_save_path,
                pdf_path=pdf_save_path,
                mode=mode,
                target_font=target_font,
                gemini_api_key=gemini_api_key
            )
            response.document_type = "pptx"
            response.filename = upload_doc.filename
            return response
    except Exception as e:
        logger.exception("Error processing document")
        raise HTTPException(status_code=500, detail=f"Failed to process document: {str(e)}")


@app.get("/api/sample", response_model=ProcessResponse)
@app.get("/sample", response_model=ProcessResponse)
def load_sample():
    """
    Generates realistic test files (corrupted PPTX and clean PDF)
    and processes them for instant 1-click testing.
    """
    session_id = str(uuid.uuid4())
    session_dir = os.path.join(BASE_TEMP_DIR, session_id)
    os.makedirs(session_dir, exist_ok=True)

    pptx_path, pdf_path = generate_sample_files(session_dir)

    sample_meta = {
        "doc_path": pptx_path,
        "pptx_path": pptx_path,
        "docx_path": None,
        "pdf_path": pdf_path,
        "original_filename": "sample_khmer_presentation.pptx",
        "target_font": "Khmer OS Battambang",
        "mode": "auto",
        "gemini_api_key": None,
        "document_type": "pptx"
    }
    save_session_meta(session_id, sample_meta)

    response = _process_presentation_internal(
        session_id=session_id,
        pptx_path=pptx_path,
        pdf_path=pdf_path,
        mode="auto",
        target_font="Khmer OS Battambang"
    )
    response.document_type = "pptx"
    response.filename = "sample_khmer_presentation.pptx"
    return response


@app.post("/api/apply-and-download")
@app.post("/apply-and-download")
def apply_and_download(payload: ApplyFixesRequest, background_tasks: BackgroundTasks):
    """
    Applies accepted replacements to the PPTX or DOCX file, sets target font,
    and returns the fixed file as a download stream.
    """
    session_id = payload.session_id
    meta = get_session_meta(session_id)
    if not meta:
        raise HTTPException(status_code=404, detail="Session expired or not found. Please upload again.")

    original_filename = meta["original_filename"]
    doc_type = meta.get("document_type", "pptx")
    
    # Filter to only accepted or modified replacements
    valid_replacements = [
        r.model_dump() for r in payload.replacements
        if r.status in ("accepted", "modified")
    ]

    target_font = payload.target_font or meta.get("target_font", "Khmer OS Battambang")

    try:
        base_name, ext = os.path.splitext(original_filename)
        fixed_filename = f"{base_name}_fixed{ext}"
        output_path = os.path.join(BASE_TEMP_DIR, session_id, fixed_filename)

        requested_format = getattr(payload, "format", "original")

        if doc_type == "pdf" and requested_format != "pptx":
            # Primary strategy: Apply fixes to the PowerPoint representation and export to PDF via Office COM
            # to guarantee native OpenType complex text shaping (proper Choeng ជើង, vowels, and accents).
            output_pptx_path = meta.get("pptx_path")
            exported_pdf = False
            if output_pptx_path and os.path.exists(output_pptx_path):
                prs = Presentation(output_pptx_path)
                apply_replacements_and_fonts(prs, valid_replacements, target_font=target_font)
                fixed_pptx_path = os.path.join(BASE_TEMP_DIR, session_id, f"{base_name}_temp_for_pdf.pptx")
                prs.save(fixed_pptx_path)
                exported_pdf = export_pptx_to_pdf(fixed_pptx_path, output_path)

            if exported_pdf and os.path.exists(output_path):
                logger.info(f"Session {session_id}: exported restored PPTX to PDF using native OpenType complex text shaping: {fixed_filename}")
                return FileResponse(
                    output_path,
                    filename=fixed_filename,
                    media_type="application/pdf"
                )

            # Fallback strategy: In-place PyMuPDF redaction/insertion
            logger.warning(f"Session {session_id}: Office PDF export unavailable. Falling back to PyMuPDF in-place restoration.")
            pdf_path = meta.get("pdf_path") or meta["doc_path"]
            apply_pdf_replacements_and_fonts(
                pdf_path=pdf_path,
                output_path=output_path,
                replacements=valid_replacements,
                target_font=target_font
            )
            return FileResponse(
                output_path,
                filename=fixed_filename,
                media_type="application/pdf"
            )
        elif doc_type == "pdf":
            converted_filename = f"{base_name}.pptx"
            output_pptx_path = meta.get("pptx_path")
            if not output_pptx_path or not os.path.exists(output_pptx_path):
                raise HTTPException(status_code=404, detail="Converted presentation not found.")
            if valid_replacements:
                prs = Presentation(output_pptx_path)
                runs_updated = apply_replacements_and_fonts(prs, valid_replacements, target_font=target_font)
                logger.info(f"Session {session_id}: applied fixes across {runs_updated} PPTX text runs from OCR.")
                fixed_pptx_path = os.path.join(BASE_TEMP_DIR, session_id, f"{base_name}_restored.pptx")
                prs.save(fixed_pptx_path)
                output_pptx_path = fixed_pptx_path
                converted_filename = f"{base_name}_restored.pptx"

            return FileResponse(
                output_pptx_path,
                filename=converted_filename,
                media_type="application/vnd.openxmlformats-officedocument.presentationml.presentation"
            )
        elif doc_type == "docx" or meta.get("docx_path"):
            docx_path = meta.get("docx_path") or meta["doc_path"]
            doc = docx.Document(docx_path)
            paragraphs_updated = apply_docx_replacements_and_fonts(
                doc=doc,
                replacements=valid_replacements,
                target_font=target_font
            )
            logger.info(f"Session {session_id}: applied fixes across {paragraphs_updated} DOCX paragraphs.")
            doc.save(output_path)

            if requested_format == "pdf":
                pdf_filename = f"{base_name}_fixed.pdf"
                pdf_output_path = os.path.join(BASE_TEMP_DIR, session_id, pdf_filename)
                if export_docx_to_pdf(output_path, pdf_output_path):
                    return FileResponse(
                        pdf_output_path,
                        filename=pdf_filename,
                        media_type="application/pdf"
                    )

            media_type = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        else:
            original_pptx_path = meta.get("pptx_path") or meta["doc_path"]
            prs = Presentation(original_pptx_path)
            runs_updated = apply_replacements_and_fonts(prs, valid_replacements, target_font=target_font)
            logger.info(f"Session {session_id}: applied fixes across {runs_updated} PPTX text runs.")
            prs.save(output_path)

            if requested_format == "pdf":
                pdf_filename = f"{base_name}_fixed.pdf"
                pdf_output_path = os.path.join(BASE_TEMP_DIR, session_id, pdf_filename)
                if export_pptx_to_pdf(output_path, pdf_output_path):
                    return FileResponse(
                        pdf_output_path,
                        filename=pdf_filename,
                        media_type="application/pdf"
                    )

            media_type = "application/vnd.openxmlformats-officedocument.presentationml.presentation"

        return FileResponse(
            output_path,
            filename=fixed_filename,
            media_type=media_type
        )
    except Exception as e:
        logger.exception("Failed to apply fixes and build document")
        raise HTTPException(status_code=500, detail=f"Failed to generate fixed document: {str(e)}")


@app.delete("/api/cleanup/{session_id}")
@app.delete("/cleanup/{session_id}")
def cleanup(session_id: str):
    """Manual cleanup endpoint."""
    cleanup_session_dir(session_id)
    return {"status": "cleaned"}


# Serve static built frontend files if available
FRONTEND_DIST = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "..", "..", "frontend", "dist"))
if os.path.exists(FRONTEND_DIST):
    from fastapi.staticfiles import StaticFiles
    app.mount("/", StaticFiles(directory=FRONTEND_DIST, html=True), name="static")

